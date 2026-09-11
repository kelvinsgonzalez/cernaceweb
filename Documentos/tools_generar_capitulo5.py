from __future__ import annotations

import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = Path("docs/generated")
IMG = OUT / "diagramas_capitulo5"
OUT.mkdir(parents=True, exist_ok=True)
IMG.mkdir(parents=True, exist_ok=True)

DOCX = OUT / "Capitulo_5_Diseno_de_la_Aplicacion_CERNACE.docx"

BLACK = "black"
WHITE = "white"


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for c in candidates:
        try:
            return ImageFont.truetype(c, size=size)
        except Exception:
            pass
    return ImageFont.load_default()


F = font(26)
FB = font(28, True)
FS = font(22)
FT = font(34, True)


def wrap(draw: ImageDraw.ImageDraw, text: str, max_width: int, fnt=F):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def text_center(draw, box, text, fnt=F):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, max(20, x2 - x1 - 24), fnt)
    line_h = fnt.size + 6
    total_h = len(lines) * line_h
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, fill=BLACK, font=fnt)
        y += line_h


def text_left(draw, xy, text, fnt=F, max_width=500, line_gap=5):
    x, y = xy
    for line in wrap(draw, text, max_width, fnt):
        draw.text((x, y), line, fill=BLACK, font=fnt)
        y += fnt.size + line_gap
    return y


def arrow(draw, start, end, width=3):
    draw.line([start, end], fill=BLACK, width=width)
    x1, y1 = start
    x2, y2 = end
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 16
    pts = [
        (x2, y2),
        (x2 - size * math.cos(ang - math.pi / 6), y2 - size * math.sin(ang - math.pi / 6)),
        (x2 - size * math.cos(ang + math.pi / 6), y2 - size * math.sin(ang + math.pi / 6)),
    ]
    draw.polygon(pts, outline=BLACK, fill=WHITE)


def rect(draw, box, text="", fnt=F, width=3):
    draw.rectangle(box, outline=BLACK, fill=WHITE, width=width)
    if text:
        text_center(draw, box, text, fnt)


def ellipse(draw, box, text="", fnt=F, width=3):
    draw.ellipse(box, outline=BLACK, fill=WHITE, width=width)
    if text:
        text_center(draw, box, text, fnt)


def actor(draw, x, y, label):
    draw.ellipse((x - 18, y, x + 18, y + 36), outline=BLACK, width=3)
    draw.line((x, y + 36, x, y + 110), fill=BLACK, width=3)
    draw.line((x - 45, y + 62, x + 45, y + 62), fill=BLACK, width=3)
    draw.line((x, y + 110, x - 45, y + 165), fill=BLACK, width=3)
    draw.line((x, y + 110, x + 45, y + 165), fill=BLACK, width=3)
    text_center(draw, (x - 100, y + 172, x + 100, y + 230), label, FS)


def save(img, name):
    path = IMG / f"{name}.png"
    img.save(path)
    return path


def diagram_use_cases():
    img = Image.new("RGB", (1800, 1100), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama de casos de uso principales", fill=BLACK, font=FT)
    rect(d, (350, 110, 1450, 1010), "", FB)
    d.rectangle((670, 118, 1130, 168), outline=BLACK, fill=WHITE, width=2)
    text_center(d, (670, 118, 1130, 168), "Sistema Web CERNACE", FS)
    actors = [("Visitante", 130, 190), ("Personal", 130, 535), ("Padrino", 1650, 360), ("Administrador", 1650, 705)]
    for label, x, y in actors:
        actor(d, x, y, label)
    cases = [
        ("Iniciar sesion", (510, 205, 850, 330)),
        ("Gestionar expediente de beneficiario", (920, 205, 1320, 330)),
        ("Registrar avance de seguimiento", (575, 440, 990, 565)),
        ("Asignar padrinazgo y consultar portal", (1030, 440, 1390, 565)),
        ("Realizar donacion o enviar solicitud", (720, 710, 1160, 840)),
    ]
    for label, box in cases:
        ellipse(d, box, label, F)
    links = [
        ((210, 310), (720, 770)), ((210, 655), (680, 500)), ((210, 655), (1080, 270)),
        ((1570, 475), (1210, 500)), ((1570, 820), (1090, 270)), ((1570, 820), (760, 265)),
        ((210, 655), (680, 265)), ((1570, 475), (680, 265)),
    ]
    for a, b in links:
        d.line([a, b], fill=BLACK, width=2)
    return save(img, "01_casos_uso")


def diagram_components():
    img = Image.new("RGB", (1800, 1250), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama de componentes", fill=BLACK, font=FT)
    rect(d, (70, 130, 400, 300), "Sitio publico\nLanding, contacto,\napadrina, donar", FS)
    rect(d, (70, 420, 400, 590), "Portal del padrino\nMis apadrinados\nAvances visibles", FS)
    rect(d, (70, 710, 400, 880), "Panel admin\nDashboard y modulos", FS)
    rect(d, (560, 140, 900, 300), "Autenticacion y RBAC\nAuth.js, JWT,\nrequirePermiso", FS)
    rect(d, (560, 380, 900, 540), "Beneficiarios y\nExpedientes", FS)
    rect(d, (560, 620, 900, 780), "Seguimiento de\nAvances", FS)
    rect(d, (560, 860, 900, 1020), "Donaciones\nPasarela simulada", FS)
    rect(d, (1080, 250, 1460, 410), "Padrinazgos\nAsignaciones", FS)
    rect(d, (1080, 520, 1460, 680), "Formularios y\nSolicitudes", FS)
    rect(d, (1080, 790, 1460, 950), "Auditoria y\nConfiguracion", FS)
    rect(d, (1530, 520, 1750, 700), "PostgreSQL\nPrisma", FS)
    arrows = [
        ((400, 215), (560, 220)), ((400, 505), (560, 220)), ((400, 795), (560, 220)),
        ((900, 220), (1080, 330)), ((900, 460), (1080, 330)), ((900, 700), (1080, 870)),
        ((900, 940), (1080, 870)), ((400, 215), (1080, 600)), ((1460, 330), (1530, 600)),
        ((900, 460), (1530, 600)), ((900, 700), (1530, 600)), ((1460, 600), (1530, 600)),
        ((1460, 870), (1530, 600)),
    ]
    for a, b in arrows:
        arrow(d, a, b)
    return save(img, "02_componentes")


def diagram_er():
    img = Image.new("RGB", (2000, 1350), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama entidad-relacion", fill=BLACK, font=FT)
    entities = {
        "User": (90, 120, 420, 290),
        "Role": (560, 120, 890, 290),
        "Permission": (1030, 120, 1360, 290),
        "Beneficiario": (90, 470, 440, 690),
        "Programa": (560, 430, 890, 590),
        "ExpedienteClinico": (1010, 410, 1390, 580),
        "FichaSocioeconomica": (1490, 410, 1900, 580),
        "Seguimiento": (1010, 690, 1390, 860),
        "Documento": (1490, 690, 1900, 860),
        "Padrino": (90, 920, 420, 1090),
        "Padrinazgo": (560, 920, 890, 1090),
        "Donacion": (1010, 980, 1390, 1150),
        "Solicitud": (1490, 980, 1900, 1150),
    }
    attrs = {
        "User": "id, nombre, email, activo",
        "Role": "id, clave, nombre",
        "Permission": "id, clave, modulo",
        "Beneficiario": "id, codigo, nombres, estado",
        "Programa": "id, nombre, activo",
        "ExpedienteClinico": "diagnostico, terapias",
        "FichaSocioeconomica": "ingreso, vulnerabilidad",
        "Seguimiento": "fecha, area, visibleParaPadrino",
        "Documento": "nombre, categoria, url",
        "Padrino": "id, nombre, email, activo",
        "Padrinazgo": "aporte, modalidad, activo",
        "Donacion": "monto, estado, referencia",
        "Solicitud": "tipo, estado, createdAt",
    }
    for name, box in entities.items():
        rect(d, box, name, FB)
        d.line((box[0], box[1] + 55, box[2], box[1] + 55), fill=BLACK, width=2)
        text_left(d, (box[0] + 18, box[1] + 75), attrs[name], FS, box[2] - box[0] - 36)
    rels = [
        ("N:M", "User", "Role"), ("N:M", "Role", "Permission"), ("1:N", "Programa", "Beneficiario"),
        ("1:1", "Beneficiario", "ExpedienteClinico"), ("1:1", "Beneficiario", "FichaSocioeconomica"),
        ("1:N", "Beneficiario", "Seguimiento"), ("1:N", "Beneficiario", "Documento"),
        ("1:N", "Padrino", "Padrinazgo"), ("1:N", "Beneficiario", "Padrinazgo"),
        ("1:N", "Padrino", "Donacion"), ("0:N", "Solicitud", "Beneficiario"),
    ]
    def center(name):
        b = entities[name]
        return ((b[0] + b[2]) // 2, (b[1] + b[3]) // 2)
    for label, a, b in rels:
        ca, cb = center(a), center(b)
        d.line([ca, cb], fill=BLACK, width=2)
        mx, my = (ca[0] + cb[0]) // 2, (ca[1] + cb[1]) // 2
        d.rectangle((mx - 38, my - 20, mx + 38, my + 20), outline=BLACK, fill=WHITE, width=1)
        text_center(d, (mx - 38, my - 20, mx + 38, my + 20), label, FS)
    return save(img, "03_entidad_relacion")


def diagram_classes():
    img = Image.new("RGB", (2000, 1350), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama de clases", fill=BLACK, font=FT)
    classes = [
        ("SesionUsuario", ["id", "email", "roles", "permisos"], ["tienePermiso()"], (80, 120, 440, 360)),
        ("User", ["id", "nombre", "email", "activo"], ["autenticar()", "actualizarAcceso()"], (560, 120, 940, 390)),
        ("Beneficiario", ["codigoExpediente", "nombres", "estado"], ["calcularEdad()", "actualizarDatos()"], (80, 520, 470, 800)),
        ("ExpedienteClinico", ["diagnostico", "terapias"], ["consultar()"], (560, 520, 940, 760)),
        ("FichaSocioeconomica", ["ingresoMensual", "nivel"], ["consultar()"], (1060, 520, 1460, 760)),
        ("Seguimiento", ["fecha", "area", "visibleParaPadrino"], ["registrar()", "publicar()"], (1540, 520, 1940, 800)),
        ("Padrino", ["nombre", "email", "activo"], ["consultarPortal()"], (80, 930, 470, 1170)),
        ("Padrinazgo", ["aporteMensual", "modalidad", "activo"], ["asignar()", "finalizar()"], (560, 930, 940, 1190)),
        ("Donacion", ["monto", "estado", "referencia"], ["iniciar()", "confirmar()"], (1060, 930, 1460, 1190)),
        ("AuditLog", ["actor", "accion", "entidad"], ["registrar()"], (1540, 930, 1940, 1170)),
    ]
    for name, attrs, methods, box in classes:
        rect(d, box, "", width=3)
        d.line((box[0], box[1] + 55, box[2], box[1] + 55), fill=BLACK, width=2)
        d.line((box[0], box[1] + 155, box[2], box[1] + 155), fill=BLACK, width=2)
        text_center(d, (box[0], box[1], box[2], box[1] + 55), name, FB)
        y = box[1] + 70
        for attr in attrs:
            d.text((box[0] + 18, y), f"+ {attr}", fill=BLACK, font=FS)
            y += 28
        y = box[1] + 170
        for m in methods:
            d.text((box[0] + 18, y), f"+ {m}", fill=BLACK, font=FS)
            y += 28
    for a, b in [((440, 250), (560, 250)), ((470, 660), (560, 640)), ((470, 660), (1060, 640)),
                 ((470, 660), (1540, 650)), ((470, 1050), (560, 1050)), ((940, 1050), (1060, 1050)),
                 ((1460, 1050), (1540, 1050))]:
        arrow(d, a, b)
    return save(img, "04_clases")


def diagram_sequence():
    img = Image.new("RGB", (2100, 1500), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagramas de secuencia - procesos principales", fill=BLACK, font=FT)
    cols = [180, 520, 860, 1200, 1540, 1880]
    labels = ["Actor", "Interfaz", "Server Action", "RBAC/Auth", "Prisma", "Base de datos"]
    for x, label in zip(cols, labels):
        rect(d, (x - 120, 120, x + 120, 180), label, FS)
        d.line((x, 180, x, 1400), fill=BLACK, width=2)
    steps = [
        ("1. Iniciar sesion", [(0,1,"credenciales"), (1,3,"validar"), (3,4,"buscar usuario"), (4,5,"consulta"), (3,1,"JWT con permisos")]),
        ("2. Gestionar expediente", [(0,1,"abrir expediente"), (1,3,"requirePermiso"), (3,4,"consultar bloques permitidos"), (4,5,"datos autorizados"), (1,0,"vista")]),
        ("3. Registrar avance", [(0,1,"formulario"), (1,2,"registrarAvance"), (2,3,"seguimiento.escribir"), (2,4,"crear Seguimiento"), (4,5,"insert")]),
        ("4. Asignar padrinazgo", [(0,1,"seleccionar padrino"), (1,2,"asignarPadrinazgo"), (2,3,"padrinazgos.gestionar"), (2,4,"crear/actualizar"), (4,5,"commit")]),
        ("5. Realizar donacion", [(0,1,"datos donacion"), (1,2,"iniciarDonacion"), (2,4,"crear Donacion"), (4,5,"pendiente"), (1,0,"pago simulado")]),
    ]
    y = 250
    for title, flow in steps:
        d.text((80, y - 42), title, fill=BLACK, font=FB)
        for a, b, msg in flow:
            arrow(d, (cols[a], y), (cols[b], y), width=2)
            text_center(d, (min(cols[a], cols[b]) + 10, y - 36, max(cols[a], cols[b]) - 10, y - 6), msg, FS)
            y += 45
        y += 85
    return save(img, "05_secuencia")


def diagram_states():
    img = Image.new("RGB", (1900, 1150), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagramas de estados", fill=BLACK, font=FT)
    tracks = [
        ("Expediente", ["INCOMPLETO", "EN_REVISION", "COMPLETO", "EGRESADO"]),
        ("Solicitud", ["NUEVA", "EN_REVISION", "APROBADA", "RECHAZADA"]),
        ("Donacion", ["PENDIENTE", "COMPLETADA", "FALLIDA", "REEMBOLSADA"]),
        ("Padrinazgo", ["SIN_ASIGNAR", "ACTIVO", "FINALIZADO"]),
        ("Avance", ["BORRADOR", "REGISTRADO", "VISIBLE_PARA_PADRINO"]),
        ("Usuario", ["ACTIVO", "SIN_ACCESO", "INACTIVO"]),
    ]
    y = 140
    for title, states in tracks:
        d.text((70, y + 30), title, fill=BLACK, font=FB)
        x = 360
        prev = None
        for state in states:
            box = (x, y, x + 250, y + 90)
            rect(d, box, state, FS)
            if prev:
                arrow(d, (prev[2], y + 45), (box[0], y + 45), width=2)
            prev = box
            x += 300
        y += 160
    return save(img, "06_estados")


def diagram_activity():
    img = Image.new("RGB", (1900, 1400), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagramas de actividades - flujo general", fill=BLACK, font=FT)
    steps = [
        "Inicio",
        "Usuario accede al sistema",
        "Validar sesion y permisos",
        "Seleccionar proceso principal",
        "Capturar o consultar informacion",
        "Validar datos del formulario",
        "Guardar cambios en base de datos",
        "Registrar auditoria",
        "Mostrar resultado al usuario",
        "Fin",
    ]
    x = 760
    y = 130
    for i, step in enumerate(steps):
        if i == 0 or i == len(steps) - 1:
            ellipse(d, (x, y, x + 360, y + 80), step, F)
        elif step == "Seleccionar proceso principal":
            pts = [(x + 180, y), (x + 360, y + 70), (x + 180, y + 140), (x, y + 70)]
            d.polygon(pts, outline=BLACK, fill=WHITE)
            text_center(d, (x + 40, y + 25, x + 320, y + 115), step, FS)
            y += 60
        else:
            rect(d, (x, y, x + 360, y + 90), step, FS)
        if i < len(steps) - 1:
            arrow(d, (x + 180, y + (140 if step == "Seleccionar proceso principal" else 90)), (x + 180, y + 135))
        y += 145
    side = [
        ("Expediente", 230, 430),
        ("Avance", 230, 575),
        ("Padrinazgo", 230, 720),
        ("Donacion", 1310, 575),
        ("Solicitud", 1310, 720),
    ]
    for label, sx, sy in side:
        rect(d, (sx, sy, sx + 280, sy + 80), label, FS)
        arrow(d, (sx + (280 if sx < 760 else 0), sy + 40), (760 if sx < 760 else 1120, 620), width=2)
    return save(img, "07_actividades")


def diagram_collaboration():
    img = Image.new("RGB", (1900, 1200), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama de colaboracion", fill=BLACK, font=FT)
    nodes = {
        "Usuario": (110, 500, 390, 620),
        "Interfaz Next.js": (560, 180, 900, 300),
        "Server Actions": (560, 500, 900, 620),
        "Auth/RBAC": (1040, 180, 1380, 300),
        "Prisma Client": (1040, 500, 1380, 620),
        "Base de datos": (1510, 500, 1810, 620),
        "Pasarela simulada": (1040, 830, 1380, 950),
        "Auditoria": (1510, 830, 1810, 950),
    }
    for label, box in nodes.items():
        rect(d, box, label, F)
    flows = [
        ("1 solicita", "Usuario", "Interfaz Next.js"),
        ("2 envia datos", "Interfaz Next.js", "Server Actions"),
        ("3 valida permiso", "Server Actions", "Auth/RBAC"),
        ("4 consulta/guarda", "Server Actions", "Prisma Client"),
        ("5 persiste", "Prisma Client", "Base de datos"),
        ("6 registra evento", "Server Actions", "Auditoria"),
        ("7 simula pago", "Server Actions", "Pasarela simulada"),
        ("8 respuesta", "Interfaz Next.js", "Usuario"),
    ]
    def mid(box):
        return ((box[0] + box[2]) // 2, (box[1] + box[3]) // 2)
    for label, a, b in flows:
        pa, pb = mid(nodes[a]), mid(nodes[b])
        arrow(d, pa, pb, width=2)
        mx, my = (pa[0] + pb[0]) // 2, (pa[1] + pb[1]) // 2
        text_center(d, (mx - 90, my - 35, mx + 90, my + 5), label, FS)
    return save(img, "08_colaboracion")


def wireframe(name: str, title: str, blocks: list[tuple[int, int, int, int, str]]):
    img = Image.new("RGB", (1600, 1000), WHITE)
    d = ImageDraw.Draw(img)
    d.text((55, 35), title, fill=BLACK, font=FT)
    rect(d, (80, 110, 1520, 930), "", width=3)
    for b in blocks:
        rect(d, b[:4], b[4], FS)
    return save(img, name)


def make_wireframes():
    return [
        wireframe("09_maqueta_login", "Maqueta - Inicio de sesion", [
            (570, 220, 1030, 300, "Logo CERNACE"), (570, 340, 1030, 410, "Campo correo"),
            (570, 435, 1030, 505, "Campo contrasena"), (690, 550, 910, 620, "Boton ingresar"),
        ]),
        wireframe("10_maqueta_panel", "Maqueta - Pantalla principal", [
            (80, 110, 330, 930, "Menu lateral"), (370, 150, 1480, 230, "Encabezado"),
            (370, 280, 620, 390, "KPI beneficiarios"), (650, 280, 900, 390, "KPI padrinos"),
            (930, 280, 1180, 390, "KPI expedientes"), (1210, 280, 1480, 390, "KPI donaciones"),
            (370, 450, 900, 820, "Ultimos avances"), (950, 450, 1480, 820, "Actividad reciente"),
        ]),
        wireframe("11_maqueta_expediente", "Maqueta - Expediente de beneficiario", [
            (100, 150, 1500, 260, "Cabecera del expediente"), (100, 300, 1000, 450, "Datos generales"),
            (100, 480, 1000, 650, "Expediente clinico"), (100, 680, 1000, 850, "Ficha socioeconomica"),
            (1060, 300, 1500, 850, "Resumen, documentos y avances"),
        ]),
        wireframe("12_maqueta_avance", "Maqueta - Registrar avance", [
            (220, 160, 1380, 240, "Beneficiario seleccionado"), (220, 290, 620, 360, "Fecha"),
            (660, 290, 1380, 360, "Area"), (220, 400, 1380, 470, "Titulo"),
            (220, 510, 1380, 720, "Descripcion"), (220, 760, 700, 830, "Visible para padrino"),
            (1160, 760, 1380, 830, "Guardar"),
        ]),
        wireframe("13_maqueta_asignacion", "Maqueta - Asignar padrinazgo", [
            (180, 160, 1420, 260, "Listado de asignaciones activas"), (180, 320, 760, 400, "Seleccionar padrino"),
            (840, 320, 1420, 400, "Seleccionar beneficiario"), (180, 455, 560, 535, "Aporte mensual"),
            (620, 455, 1000, 535, "Modalidad"), (1060, 455, 1420, 535, "Fecha inicio"),
            (1180, 600, 1420, 680, "Asignar"),
        ]),
        wireframe("14_maqueta_portal", "Maqueta - Portal del padrino", [
            (140, 160, 1460, 250, "Saludo y resumen"), (140, 310, 700, 720, "Tarjeta apadrinado"),
            (760, 310, 1460, 720, "Avances publicados"), (1160, 790, 1460, 860, "Ver progreso"),
        ]),
        wireframe("15_maqueta_donacion", "Maqueta - Donacion", [
            (190, 160, 700, 790, "Formulario de donante"), (780, 160, 1410, 360, "Campanas activas"),
            (780, 410, 1410, 620, "Resumen de aporte"), (1160, 700, 1410, 780, "Continuar al pago"),
        ]),
        wireframe("16_maqueta_solicitudes", "Maqueta - Solicitudes", [
            (120, 150, 1480, 240, "Filtros y contadores"), (120, 300, 1480, 760, "Tabla de solicitudes"),
            (1080, 820, 1480, 890, "Cambiar estado"),
        ]),
    ]


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")
        borders.append(tag)
    tblPr.append(borders)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    set_borders(table)
    doc.add_paragraph()
    return table


def add_picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(6.35))
    picture._inline.docPr.set("title", caption.split(". ", 1)[-1])
    picture._inline.docPr.set("descr", caption)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9)
    return p


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 14)]:
        st = styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = None
        st.font.bold = False
        st.paragraph_format.space_before = Pt(14)
        st.paragraph_format.space_after = Pt(6)


def add_fill_space(doc, lines=4):
    p = doc.add_paragraph("Espacio para completar:")
    p.runs[0].bold = True
    for _ in range(lines):
        doc.add_paragraph("_" * 86)


def add_function_prototype(doc, titulo, ruta, descripcion, validaciones):
    doc.add_heading(titulo, level=3)
    doc.add_paragraph(descripcion)
    add_table(
        doc,
        ["Pantalla/Ruta", "Criterios de validación"],
        [[ruta, "; ".join(validaciones)]],
    )


def build_doc():
    paths = [
        diagram_use_cases(),
        diagram_components(),
        diagram_er(),
        diagram_classes(),
        diagram_sequence(),
        diagram_states(),
        diagram_activity(),
        diagram_collaboration(),
        *make_wireframes(),
    ]

    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("CAPÍTULO 5\nDISEÑO DE LA APLICACIÓN")
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.bold = True
    doc.add_paragraph("Proyecto: Plataforma web CERNACE").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Documento base para completar el diseño de la aplicación, con diagramas monocromos y terminología uniforme.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("5.1. Desarrollo de la Propuesta", level=1)
    doc.add_paragraph(
        "La propuesta se organiza en seis módulos principales: Autenticación y RBAC, "
        "Beneficiarios y Expedientes, Seguimiento de Avances, Padrinazgos y Portal, "
        "Donaciones, y Formularios/Solicitudes. Estos módulos cubren los procesos "
        "centrales del sistema y se conectan con la base de datos mediante Prisma."
    )
    add_table(
        doc,
        ["Orden", "Módulo principal", "Responsabilidad"],
        [
            ["1", "Autenticación y RBAC", "Inicio de sesión, roles, permisos y control de acceso."],
            ["2", "Beneficiarios y Expedientes", "Datos generales, clínicos, socioeconómicos y documentos."],
            ["3", "Seguimiento de Avances", "Registro de avances y visibilidad para padrinos."],
            ["4", "Padrinazgos y Portal", "Asignación padrino-beneficiario y consulta privada."],
            ["5", "Donaciones", "Registro de aportes y pasarela simulada."],
            ["6", "Formularios/Solicitudes", "Inscripciones, contacto y revisión administrativa."],
        ],
    )

    doc.add_page_break()
    doc.add_heading("5.1.1. Diagramas de casos de uso", level=2)
    doc.add_paragraph("Se consideran cinco casos de uso principales, respetando el límite solicitado.")
    add_picture(doc, paths[0], "Figura 5.1. Casos de uso principales del sistema.")
    add_table(
        doc,
        ["Caso de uso", "Actor principal", "Resultado esperado"],
        [
            ["Iniciar sesión", "Personal, administrador o padrino", "Sesión activa con permisos cargados."],
            ["Gestionar expediente", "Personal autorizado", "Expediente consultado o actualizado."],
            ["Registrar avance", "Terapeuta o trabajo social", "Avance guardado y auditado."],
            ["Asignar padrinazgo y consultar portal", "Administrador, padrino", "Asignación activa y consulta segura."],
            ["Realizar donación o enviar solicitud", "Visitante", "Donación o solicitud registrada."],
        ],
    )

    doc.add_page_break()
    doc.add_heading("5.1.2. Diagramas de componentes", level=2)
    add_picture(doc, paths[1], "Figura 5.2. Componentes principales de la aplicación.")

    doc.add_page_break()
    doc.add_heading("5.1.3. Diagrama de modelo entidad-relación", level=2)
    add_picture(doc, paths[2], "Figura 5.3. Modelo entidad-relación simplificado.")

    doc.add_page_break()
    doc.add_heading("5.1.4. Diagrama de clases", level=2)
    add_picture(doc, paths[3], "Figura 5.4. Clases conceptuales del sistema.")

    doc.add_page_break()
    doc.add_heading("5.1.5. Diagramas de secuencia", level=2)
    add_picture(doc, paths[4], "Figura 5.5. Secuencias de los cinco procesos principales.")

    doc.add_page_break()
    doc.add_heading("5.1.6. Diagramas de estados", level=2)
    add_picture(doc, paths[5], "Figura 5.6. Estados principales por módulo.")

    doc.add_page_break()
    doc.add_heading("5.1.7. Diagramas de actividades", level=2)
    add_picture(doc, paths[6], "Figura 5.7. Actividades generales del sistema.")

    doc.add_page_break()
    doc.add_heading("5.1.8. Diagramas de colaboración", level=2)
    add_picture(doc, paths[7], "Figura 5.8. Colaboración entre interfaz, servidor, RBAC, Prisma y base de datos.")

    doc.add_page_break()
    doc.add_heading("5.2. Maquetización", level=1)
    doc.add_paragraph(
        "Las maquetas se presentan en blanco y negro, con distribución funcional. "
        "Su finalidad es definir ubicación de controles, jerarquía visual y navegación principal."
    )
    wire_titles = [
        ("5.2.1. Maqueta de Inicio de Sesión", 8, "Figura 5.9"),
        ("5.2.2. Maqueta de la pantalla principal", 9),
        ("5.2.3.1. Maqueta principal: Expediente de beneficiario", 10, "Figura 5.11"),
        ("5.2.3.2. Maqueta principal: Registrar avance", 11, "Figura 5.12"),
        ("5.2.3.3. Maqueta principal: Asignar padrinazgo", 12, "Figura 5.13"),
        ("5.2.3.4. Maqueta principal: Portal del padrino", 13, "Figura 5.14"),
        ("5.2.3.5. Maqueta principal: Donación", 14, "Figura 5.15"),
        ("5.2.3.6. Maqueta principal: Solicitudes", 15, "Figura 5.16"),
    ]
    wire_titles[1] = (wire_titles[1][0], wire_titles[1][1], "Figura 5.10")
    for heading, idx, figura in wire_titles:
        doc.add_page_break()
        doc.add_heading(heading, level=2)
        add_picture(doc, paths[idx], f"{figura}. {heading}.")
        add_fill_space(doc, 2)

    doc.add_page_break()
    doc.add_heading("5.3. Prototipo", level=1)
    doc.add_paragraph(
        "El prototipo funcional corresponde a las pantallas implementadas en la aplicación Next.js. "
        "Las siguientes fichas describen la función esperada y el criterio mínimo de validación."
    )
    doc.add_heading("5.3.1. Prototipo de la pantalla principal", level=2)
    add_table(
        doc,
        ["Elemento", "Función esperada", "Validación"],
        [
            ["Menú lateral", "Acceso a módulos según permisos.", "El usuario solo ve rutas autorizadas."],
            ["Indicadores KPI", "Mostrar resumen de beneficiarios, padrinos y expedientes.", "Los conteos coinciden con la base de datos."],
            ["Actividad reciente", "Mostrar avances o auditoría según rol.", "No se consultan datos sin permiso."],
        ],
    )
    add_fill_space(doc, 3)

    doc.add_heading("5.3.2. Prototipo de las distintas funciones principales", level=2)
    add_table(
        doc,
        ["Función principal", "Pantalla/Ruta", "Criterio de aceptación"],
        [
            ["Iniciar sesión", "/login", "Credenciales válidas crean sesión JWT con roles y permisos."],
            ["Gestionar expediente", "/admin/beneficiarios/[id]", "El rol visualiza solo los bloques autorizados."],
            ["Registrar avance", "/admin/beneficiarios/[id]/avance", "El avance se guarda y puede marcarse visible para padrino."],
            ["Asignar padrinazgo", "/admin/asignaciones", "Solo un padrinazgo activo por beneficiario."],
            ["Consultar portal", "/portal/[id]", "El padrino solo ve sus apadrinados y avances publicados."],
            ["Realizar donación", "/donar", "La donación se registra con referencia y estado."],
        ],
    )
    prototipos = [
        (
            "Inicio de sesión",
            "/login",
            "Permite que personal, administradores y padrinos ingresen con credenciales y reciban una sesión con permisos.",
            ["valida correo y contraseña", "redirige según rol", "bloquea cuentas inactivas"],
        ),
        (
            "Gestión de expediente",
            "/admin/beneficiarios/[id]",
            "Muestra datos generales, clínicos, socioeconómicos, documentos, citas y avances del beneficiario según permisos.",
            ["oculta bloques no autorizados", "registra auditoría de apertura", "mantiene navegación al listado"],
        ),
        (
            "Registro de avance",
            "/admin/beneficiarios/[id]/avance",
            "Captura avances por área de atención y permite definir si el padrino podrá consultarlos en su portal.",
            ["guarda fecha, área, título y descripción", "respeta permiso de seguimiento", "actualiza el expediente"],
        ),
        (
            "Asignación de padrinazgo",
            "/admin/asignaciones",
            "Vincula un padrino con un beneficiario activo que aún no cuenta con apoyo asignado.",
            ["evita duplicar padrinazgo activo", "permite finalizar asignación", "actualiza galería pública y portal"],
        ),
        (
            "Portal del padrino",
            "/portal/[id]",
            "Permite al padrino consultar únicamente sus beneficiarios asignados y los avances marcados como visibles.",
            ["filtra por padrino de la sesión", "devuelve 404 para beneficiarios ajenos", "muestra solo avances públicos"],
        ),
        (
            "Donación pública",
            "/donar",
            "Registra aportes únicos o recurrentes mediante una pasarela simulada y deja comprobante de resultado.",
            ["crea referencia de pasarela", "actualiza estado de pago", "registra auditoría de la transacción"],
        ),
    ]
    for item in prototipos:
        add_function_prototype(doc, *item)
    add_fill_space(doc, 6)

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build_doc()
